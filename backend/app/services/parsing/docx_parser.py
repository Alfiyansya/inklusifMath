"""
DOCX Document Parser.

Extracts document structure (headings, paragraphs, tables) and
math expressions (OMML → LaTeX) from .docx files using:
  - python-docx: document structure and text
  - zipfile + xml.etree.ElementTree: OMML math nodes
  - omml_to_latex: converts OMML XML to LaTeX strings

Pipeline (per TDD Section 11.1):
  DOCX file
    │
    ├─ python-docx → Extract headings, paragraphs, tables
    │
    ├─ zipfile + ElementTree → Extract <m:oMath> XML nodes
    │     │
    │     └─ omml_to_latex converter → LaTeX strings
    │
    └─ Merge: structured content + LaTeX expressions
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Generator

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.services.parsing.models import (
    BlockType,
    ContentBlock,
    MathExpressionResult,
    ParsedDocument,
)
from app.services.parsing.omml_to_latex import omml_element_to_latex

logger = logging.getLogger(__name__)

# OMML namespace
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MATH_TAG = f"{{{_MATH_NS}}}oMath"

# Regex to detect inline math patterns in plain text (fallback)
_INLINE_MATH_PATTERNS = [
    re.compile(r"\$\$(.+?)\$\$", re.DOTALL),    # $$...$$
    re.compile(r"\$(.+?)\$"),                     # $...$
    re.compile(r"\\begin\{.+?\}.*?\\end\{.+?\}", re.DOTALL),  # LaTeX environments
]


class DocxParser:
    """
    Parses a .docx file to extract structured content and math expressions.

    Usage:
        parser = DocxParser()
        result: ParsedDocument = parser.parse(file_path_or_bytes)
    """

    def parse(self, source: str | Path | bytes, title: str = "") -> ParsedDocument:
        """
        Parse a DOCX file.

        Args:
            source: File path (str/Path) or raw bytes of the .docx file.
            title: Human-readable document title (falls back to filename or first heading).

        Returns:
            ParsedDocument with content blocks and math expressions.
        """
        try:
            if isinstance(source, (str, Path)):
                file_bytes = Path(source).read_bytes()
                if not title:
                    title = Path(source).stem
            else:
                file_bytes = source

            return self._parse_bytes(file_bytes, title)

        except Exception as exc:
            logger.error("DOCX parsing failed: %s", exc, exc_info=True)
            return ParsedDocument(
                title=title or "Unknown",
                file_type="docx",
                parse_error=str(exc),
            )

    def _parse_bytes(self, file_bytes: bytes, title: str) -> ParsedDocument:
        """Internal: parse raw .docx bytes."""
        # Load via python-docx for structure
        doc = docx.Document(io.BytesIO(file_bytes))

        # Extract OMML math expressions from the raw XML (zipfile layer)
        math_lookup = self._extract_omml_math(file_bytes)

        content_blocks: list[ContentBlock] = []
        math_expressions: list[MathExpressionResult] = []
        math_counter = 0

        for block_idx, block in enumerate(self._iter_blocks(doc)):
            if isinstance(block, Paragraph):
                para_result = self._process_paragraph(
                    block, block_idx, math_lookup, math_counter
                )
                if para_result:
                    content_block, new_math = para_result
                    content_blocks.append(content_block)
                    math_expressions.extend(new_math)
                    math_counter += len(new_math)
            elif isinstance(block, Table):
                table_block = self._process_table(block, block_idx)
                if table_block:
                    content_blocks.append(table_block)

        # Derive title from first heading if not provided
        if not title:
            for blk in content_blocks:
                if blk.block_type == BlockType.HEADING1:
                    title = blk.text.strip()
                    break

        parsed = ParsedDocument(
            title=title or "Dokumen Tanpa Judul",
            file_type="docx",
            content_blocks=content_blocks,
            math_expressions=math_expressions,
        )
        parsed.raw_structure = parsed.to_raw_structure()
        return parsed

    # ── OMML Extraction ──────────────────────────────────────────────────────

    def _extract_omml_math(self, file_bytes: bytes) -> list[str]:
        """
        Extract all OMML <m:oMath> elements from document.xml inside the DOCX zip.

        Returns a list of LaTeX strings in document order.
        """
        latex_list: list[str] = []
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                if "word/document.xml" not in zf.namelist():
                    return latex_list

                xml_bytes = zf.read("word/document.xml")
                root = ET.fromstring(xml_bytes)

                for math_elem in root.iter(_MATH_TAG):
                    latex = omml_element_to_latex(math_elem)
                    # Fallback: use original notation text if converter returns empty
                    if not latex:
                        latex = self._extract_math_text_fallback(math_elem)
                    latex_list.append(latex)

        except (zipfile.BadZipFile, ET.ParseError) as exc:
            logger.warning("Failed to extract OMML from zip: %s", exc)

        return latex_list

    def _extract_math_text_fallback(self, math_elem: ET.Element) -> str:
        """Extract plain text from an OMML element as fallback."""
        texts = []
        for t in math_elem.iter(f"{{{_MATH_NS}}}t"):
            if t.text:
                texts.append(t.text)
        return "".join(texts)

    # ── Document Structure ───────────────────────────────────────────────────

    def _iter_blocks(self, doc: docx.Document) -> Generator:
        """
        Iterate over top-level blocks (paragraphs and tables) in document order.
        This is necessary because doc.paragraphs skips tables.
        """
        for child in doc.element.body:
            tag = child.tag
            if tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif tag == qn("w:tbl"):
                yield Table(child, doc)

    def _process_paragraph(
        self,
        para: Paragraph,
        block_idx: int,
        math_lookup: list[str],
        math_start: int,
    ) -> tuple[ContentBlock, list[MathExpressionResult]] | None:
        """
        Process a single paragraph.

        Detects:
        - Headings (style name starts with 'Heading')
        - Math expressions embedded as OMML runs
        - Plain text paragraphs

        Returns (ContentBlock, list of MathExpressionResult) or None if empty.
        """
        # Determine block type from style
        block_type, level = self._get_block_type(para)

        # Collect OMML math nodes within this paragraph
        para_math_elems = list(para._element.iter(_MATH_TAG))

        # Build text representation with math placeholders
        full_text = self._get_paragraph_text_with_placeholders(
            para, len(para_math_elems), math_start
        )

        if not full_text.strip() and not para_math_elems:
            return None  # skip empty paragraphs

        content_block = ContentBlock(
            block_type=block_type,
            text=full_text,
            level=level,
            index=block_idx,
        )

        # Build MathExpressionResult for each OMML math in this paragraph
        math_results: list[MathExpressionResult] = []
        for local_idx, math_elem in enumerate(para_math_elems):
            global_idx = math_start + local_idx
            # Use the pre-extracted latex from zip parse if available
            if global_idx < len(math_lookup):
                latex = math_lookup[global_idx]
            else:
                latex = omml_element_to_latex(math_elem)

            original_notation = self._extract_math_text_fallback(math_elem)
            confidence = self._estimate_confidence(latex, original_notation)

            math_results.append(
                MathExpressionResult(
                    original_notation=original_notation or latex,
                    latex_representation=latex,
                    position_order=global_idx + 1,  # 1-based
                    block_index=block_idx,
                    conversion_confidence=confidence,
                )
            )

        return content_block, math_results

    def _process_table(self, table: Table, block_idx: int) -> ContentBlock | None:
        """Convert a Word table to a plain-text representation."""
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(cells)

        if not rows:
            return None

        # Render as pipe-separated text (Markdown-like)
        lines = []
        for i, row in enumerate(rows):
            lines.append(" | ".join(row))
            if i == 0:
                lines.append(" | ".join("---" for _ in row))  # header separator

        return ContentBlock(
            block_type=BlockType.TABLE,
            text="\n".join(lines),
            level=0,
            index=block_idx,
        )

    # ── Helpers ──────────────────────────────────────────────────────────────

    def _get_block_type(self, para: Paragraph) -> tuple[BlockType, int]:
        """Determine BlockType and heading level from paragraph style."""
        style_name = para.style.name if para.style else ""

        if style_name.startswith("Heading 1") or style_name == "Title":
            return BlockType.HEADING1, 1
        if style_name.startswith("Heading 2"):
            return BlockType.HEADING2, 2
        if style_name.startswith("Heading 3"):
            return BlockType.HEADING3, 3
        if style_name.startswith("Heading"):
            # Generic heading beyond level 3 — treat as H3
            return BlockType.HEADING3, 3

        return BlockType.PARAGRAPH, 0

    def _get_paragraph_text_with_placeholders(
        self,
        para: Paragraph,
        math_count: int,
        math_start: int,
    ) -> str:
        """
        Build the full paragraph text, replacing OMML math elements
        with readable placeholders like {{MATH_1}}.
        """
        # Get text from python-docx runs
        text_parts: list[str] = []
        local_math = 0

        for elem in para._element:
            tag = elem.tag

            if tag == qn("w:r"):
                # Normal text run
                run_text = "".join(
                    t.text or ""
                    for t in elem.iter(qn("w:t"))
                )
                text_parts.append(run_text)

            elif tag == _MATH_TAG:
                # Math block at paragraph level
                placeholder = f"{{{{MATH_{math_start + local_math + 1}}}}}"
                text_parts.append(f" {placeholder} ")
                local_math += 1

            elif tag == qn("w:ins") or tag == qn("w:hyperlink"):
                # Tracked changes / hyperlinks — recurse for text
                for sub_run in elem.iter(qn("w:r")):
                    run_text = "".join(
                        t.text or "" for t in sub_run.iter(qn("w:t"))
                    )
                    text_parts.append(run_text)

        return "".join(text_parts)

    def _estimate_confidence(self, latex: str, original: str) -> float:
        """
        Estimate conversion confidence score (0.0–1.0).

        Rules:
        - Empty latex → 0.0
        - latex == original (no conversion) → 0.5
        - Contains LaTeX commands → 0.95
        - Simple alphanumeric → 0.9
        """
        if not latex:
            return 0.0
        if latex == original:
            return 0.5
        if "\\" in latex:
            return 0.95
        if re.match(r"^[a-zA-Z0-9\s\+\-\=\*\.]+$", latex):
            return 0.9
        return 0.85
